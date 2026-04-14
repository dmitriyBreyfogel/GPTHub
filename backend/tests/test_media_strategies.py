from __future__ import annotations

import asyncio
import json
import sys
from io import BytesIO
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch


BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from fastapi import FastAPI
from fastapi.testclient import TestClient
from pptx import Presentation

from app.api.v1 import files as files_module
from app.core.config import settings
from app.core.prompt_cache import prompt_cache_manager
from app.storage.files import FileQuotaExceededError, StoredFile, build_file_access_token
from app.strategies.base import StrategyRequest, StrategyResponse, TaskType
from app.strategies.audio import AudioStrategy
from app.strategies.image_gen import ImageGenStrategy
from app.strategies.presentation import PresentationStrategy, SlideSpec
from app.strategies.response_utils import stream_chunk
from app.strategies.vision import VisionStrategy


class _FakeImageClient:
    def __init__(
        self,
        urls_by_model: dict[str, str],
        failing_models: set[str] | None = None,
    ) -> None:
        self.urls_by_model = urls_by_model
        self.failing_models = failing_models or set()
        self.calls: list[dict] = []

    async def generate_image(self, prompt: str, *, model: str | None = None) -> str:
        self.calls.append({"prompt": prompt, "model": model})
        if model in self.failing_models:
            raise RuntimeError(f"{model} is unavailable")
        return self.urls_by_model[model or settings.image_generation_model]


class _FakePresentationClient:
    def __init__(self, response: dict) -> None:
        self.response = response
        self.chat_calls: list[dict] = []
        self.image_calls: list[dict] = []

    async def chat(
        self,
        messages,
        model: str | None = None,
        temperature: float = 0.7,
        generation_options: dict | None = None,
    ):
        self.chat_calls.append(
            {
                "messages": messages,
                "model": model,
                "temperature": temperature,
                "generation_options": generation_options,
            }
        )
        return type("ChatResponse", (), {"content": json.dumps(self.response, ensure_ascii=False)})()

    async def generate_image(self, prompt: str, model: str | None = None) -> str:
        self.image_calls.append({"prompt": prompt, "model": model})
        return "https://cdn.example/presentation.png"


class _FakeFileStorage:
    def __init__(self) -> None:
        self.upload_calls: list[dict] = []
        self.download_calls: list[dict] = []
        self.shared_download_calls: list[dict] = []

    async def upload(
        self,
        *,
        file_bytes: bytes,
        filename: str,
        content_type: str,
        user_id: str,
    ) -> StoredFile:
        self.upload_calls.append(
            {
                "file_bytes": file_bytes,
                "filename": filename,
                "content_type": content_type,
                "user_id": user_id,
            }
        )
        return StoredFile(
            file_id="presentation-file-id",
            bucket="gpthub",
            object_key="anonymous/presentation-file-id/deck.pptx",
            content_type=content_type,
            size_bytes=len(file_bytes),
        )

    async def download(self, *, file_id: str, user_id: str) -> tuple[bytes, str]:
        self.download_calls.append({"file_id": file_id, "user_id": user_id})
        return b"pptx-bytes", "application/vnd.openxmlformats-officedocument.presentationml.presentation"

    async def download_shared(self, *, file_id: str) -> tuple[bytes, str]:
        self.shared_download_calls.append({"file_id": file_id})
        return b"pptx-bytes", "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _strategy_request(
    *,
    task_type: TaskType,
    text: str,
    model_override: str | None = None,
) -> StrategyRequest:
    return StrategyRequest(
        task_type=task_type,
        text=text,
        user_id="anonymous",
        model_override=model_override,
    )


class AudioStrategyTests(unittest.IsolatedAsyncioTestCase):
    async def test_stream_emits_heartbeat_without_restarting_audio_processing(self) -> None:
        strategy = AudioStrategy()
        request = StrategyRequest(
            task_type=TaskType.AUDIO,
            text="summarize",
            user_id="user-1",
            model_override="asr-model",
            file_bytes=b"audio-bytes",
            file_name="song.mp3",
            file_content_type="audio/mpeg",
        )

        async def slow_transcribe(_request):
            await asyncio.sleep(0.01)
            return "audio transcript"

        async def stream_text(_request):
            yield b"data: audio answer\n\n"
            yield b"data: [DONE]\n\n"

        transcribe = AsyncMock(side_effect=slow_transcribe)

        with (
            patch.object(strategy, "_transcribe", transcribe),
            patch.object(strategy._text_strategy, "stream", stream_text),
            patch("app.strategies.audio.AUDIO_STREAM_HEARTBEAT_SECONDS", 0.001),
        ):
            chunks = [chunk async for chunk in strategy.stream(request)]

        transcribe.assert_awaited_once_with(request)
        first_payload = json.loads(chunks[0].decode("utf-8").removeprefix("data: "))
        self.assertEqual({}, first_payload["choices"][0]["delta"])
        self.assertIn("audio answer", b"".join(chunks).decode("utf-8"))
        self.assertEqual(b"data: [DONE]\n\n", chunks[-1])

    async def test_execute_uses_routing_text_model_for_final_synthesis(self) -> None:
        strategy = AudioStrategy()
        request = StrategyRequest(
            task_type=TaskType.AUDIO,
            text="summarize",
            user_id="user-1",
            model_override="custom-audio-model",
            request_files=[
                type("AudioFile", (), {
                    "file_bytes": b"audio-bytes",
                    "file_name": "song.wav",
                    "file_content_type": "audio/wav",
                    "file_url": None,
                })()
            ],
            routing_models={"text": "custom-text-model"},
        )

        with (
            patch("app.strategies.audio.mws_client.transcribe", new=AsyncMock(return_value="hello world")) as transcribe,
            patch.object(strategy._text_strategy, "execute", new=AsyncMock(return_value=StrategyResponse(
                content="done",
                model_used="custom-text-model",
                task_type=TaskType.TEXT,
                routing_reason="text",
            ))) as text_execute,
        ):
            response = await strategy.execute(request)

        self.assertEqual("done", response.content)
        self.assertEqual("custom-text-model", response.model_used)
        transcribe.assert_awaited_once()
        forwarded_request = text_execute.await_args.args[0]
        self.assertEqual("custom-text-model", forwarded_request.model_override)

    async def test_transcribe_concatenates_multiple_audio_attachments(self) -> None:
        strategy = AudioStrategy()
        request = StrategyRequest(
            task_type=TaskType.AUDIO,
            text="analyze",
            user_id="user-1",
            model_override="custom-audio-model",
            request_files=[
                type("AudioFile", (), {
                    "file_bytes": b"one",
                    "file_name": "one.wav",
                    "file_content_type": "audio/wav",
                    "file_url": None,
                })(),
                type("AudioFile", (), {
                    "file_bytes": b"two",
                    "file_name": "two.wav",
                    "file_content_type": "audio/wav",
                    "file_url": None,
                })(),
            ],
        )

        with patch("app.strategies.audio.mws_client.transcribe", new=AsyncMock(side_effect=["first", "second"])) as transcribe:
            transcript = await strategy._transcribe(request)

        self.assertIn("[one.wav]", transcript)
        self.assertIn("[two.wav]", transcript)
        self.assertEqual(2, transcribe.await_count)


class VisionStrategyTests(unittest.IsolatedAsyncioTestCase):
    async def test_build_messages_includes_multiple_images_and_audio_context(self) -> None:
        strategy = VisionStrategy()
        request = StrategyRequest(
            task_type=TaskType.IMAGE_ANALYSIS,
            text="compare the images",
            user_id="user-1",
            model_override="custom-vision-model",
            request_files=[
                type("ImageFile", (), {
                    "file_bytes": b"image-one",
                    "file_name": "one.png",
                    "file_content_type": "image/png",
                    "file_url": None,
                })(),
                type("ImageFile", (), {
                    "file_bytes": b"image-two",
                    "file_name": "two.png",
                    "file_content_type": "image/png",
                    "file_url": None,
                })(),
                type("AudioFile", (), {
                    "file_bytes": b"audio",
                    "file_name": "note.wav",
                    "file_content_type": "audio/wav",
                    "file_url": None,
                })(),
            ],
            context_messages=[{"role": "user", "content": "compare the images"}],
        )

        with patch("app.strategies.vision.mws_client.transcribe", new=AsyncMock(return_value="spoken hint")):
            messages = await strategy._build_messages(request)

        self.assertEqual("system", messages[0].role)
        self.assertEqual("user", messages[1].role)
        content = messages[1].content
        self.assertIsInstance(content, list)
        image_parts = [item for item in content if item.get("type") == "image_url"]
        self.assertEqual(2, len(image_parts))
        text_parts = [item for item in content if item.get("type") == "text"]
        self.assertTrue(any("spoken hint" in item.get("text", "") for item in text_parts))


class ImageGenerationStrategyTests(unittest.IsolatedAsyncioTestCase):
    async def test_execute_returns_ready_markdown_image_instead_of_background_task(self) -> None:
        strategy = ImageGenStrategy()
        image_url = "https://cdn.example/generated.png"
        fake_client = _FakeImageClient({settings.image_generation_model: image_url})

        with patch("app.strategies.image_gen.mws_client", fake_client):
            response = await strategy.execute(
                _strategy_request(
                    task_type=TaskType.IMAGE_GEN,
                    text="Нарисуй красный трамвай ночью",
                    model_override=settings.image_generation_model,
                )
            )

        self.assertEqual(settings.image_generation_model, response.model_used)
        self.assertEqual(image_url, response.image_url)
        self.assertEqual([image_url], response.sources)
        self.assertIsNone(response.task_id)
        self.assertIn(f"![Сгенерированное изображение]({image_url})", response.content)
        self.assertEqual(
            [{"prompt": "Нарисуй красный трамвай ночью", "model": settings.image_generation_model}],
            fake_client.calls,
        )

    async def test_execute_uses_fallback_image_model_when_primary_fails(self) -> None:
        strategy = ImageGenStrategy()
        image_url = "https://cdn.example/fallback.png"
        fake_client = _FakeImageClient(
            {settings.image_generation_fallback_model: image_url},
            failing_models={settings.image_generation_model},
        )

        with patch("app.strategies.image_gen.mws_client", fake_client):
            response = await strategy.execute(
                _strategy_request(
                    task_type=TaskType.IMAGE_GEN,
                    text="Generate a watercolor city skyline",
                    model_override=settings.image_generation_model,
                )
            )

        self.assertEqual(settings.image_generation_fallback_model, response.model_used)
        self.assertEqual(image_url, response.image_url)
        self.assertIn("Резервная модель", response.content)
        self.assertEqual(
            [
                {"prompt": "Generate a watercolor city skyline", "model": settings.image_generation_model},
                {"prompt": "Generate a watercolor city skyline", "model": settings.image_generation_fallback_model},
            ],
            fake_client.calls,
        )

    async def test_stream_includes_generated_image_metadata(self) -> None:
        strategy = ImageGenStrategy()
        image_url = "https://cdn.example/streamed.png"
        fake_client = _FakeImageClient({settings.image_generation_model: image_url})

        with patch("app.strategies.image_gen.mws_client", fake_client):
            chunks = [
                chunk
                async for chunk in strategy.stream(
                    _strategy_request(
                        task_type=TaskType.IMAGE_GEN,
                        text="Нарисуй постер",
                        model_override=settings.image_generation_model,
                    )
                )
            ]

        first_payload = json.loads(chunks[0].decode("utf-8").removeprefix("data: "))
        self.assertEqual(image_url, first_payload["gpthub"]["image_url"])
        self.assertEqual("image_gen", first_payload["gpthub"]["task_type"])
        self.assertIn(
            f"![Сгенерированное изображение]({image_url})",
            first_payload["choices"][0]["delta"]["content"],
        )


class PresentationStrategyTests(unittest.IsolatedAsyncioTestCase):
    def test_presentation_prompt_requests_visual_slide_contract(self) -> None:
        prompt = prompt_cache_manager.build_presentation_system_prompt()

        for marker in (
            '"takeaway"',
            '"visual_hint"',
            '"layout"',
            '"body"',
            '"image_prompt"',
            "comparison",
            "timeline",
            "metrics",
            "statement",
            "image_text",
            "process",
        ):
            self.assertIn(marker, prompt)
        self.assertIn("6-10", prompt)
        self.assertIn("1-2", prompt)
        self.assertIn("Сначала проектируй содержание", prompt)
        self.assertIn("дошкольных", prompt)
        self.assertIn("не фоновую сцену", prompt)

    async def test_execute_returns_markdown_download_link(self) -> None:
        strategy = PresentationStrategy()
        fake_storage = _FakeFileStorage()

        with (
            patch.object(
                strategy,
                "_generate_slide_specs",
                new=AsyncMock(
                    return_value=[
                        SlideSpec(title="Архитектура", bullets=["Backend", "Frontend"]),
                    ]
                ),
            ),
            patch.object(strategy, "_build_pptx", return_value=b"pptx-bytes"),
            patch("app.strategies.presentation.file_storage", fake_storage),
        ):
            response = await strategy.execute(
                _strategy_request(
                    task_type=TaskType.PRESENTATION,
                    text="Сделай презентацию про архитектуру",
                    model_override="slides-model",
                )
            )

        expected_url = f"/v1/files/presentation-file-id?access_token={build_file_access_token('presentation-file-id')}"
        self.assertEqual("slides-model", response.model_used)
        self.assertEqual(expected_url, response.file_url)
        self.assertIn(f"[Скачать презентацию]({expected_url})", response.content)
        self.assertEqual("anonymous", fake_storage.upload_calls[0]["user_id"])

    def test_parse_slides_accepts_design_fields(self) -> None:
        strategy = PresentationStrategy()

        slides = strategy._parse_slides(
            {
                "slides": [
                    {
                        "title": "Архитектура GPTHub",
                        "subtitle": "Поток запроса от UI до модели",
                        "bullets": ["Frontend", "Router", "Strategy", "MWS GPT"],
                        "takeaway": "Маршрутизация отделяет UX от выбора модели.",
                        "visual_hint": "Схема из четырех связанных блоков",
                        "layout": "two_column",
                        "speaker_notes": "Пояснить, где принимается решение о task type.",
                    },
                    {
                        "title": "Сценарий в действии",
                        "layout": "image_text",
                        "body": "Пользователь видит единый поток вместо набора разрозненных инструментов.",
                        "bullets": ["Единый чат", "Автоматический выбор", "Понятный результат"],
                        "left_title": "До",
                        "right_title": "После",
                        "left_items": ["Ручной выбор сервиса", "Потеря контекста"],
                        "right_items": ["Автоматическая маршрутизация", "Сохранение памяти"],
                        "metrics": [{"value": "1", "label": "единый интерфейс"}],
                        "image_prompt": "Сотрудник работает в едином корпоративном AI-чате",
                    }
                ]
            }
        )

        self.assertEqual(2, len(slides))
        self.assertEqual("title", slides[0].layout)
        self.assertEqual("Поток запроса от UI до модели", slides[0].subtitle)
        self.assertEqual("Маршрутизация отделяет UX от выбора модели.", slides[0].takeaway)
        self.assertEqual("Схема из четырех связанных блоков", slides[0].visual_hint)
        self.assertEqual("image_text", slides[1].layout)
        self.assertEqual("Пользователь видит единый поток вместо набора разрозненных инструментов.", slides[1].body)
        self.assertEqual(["Ручной выбор сервиса", "Потеря контекста"], slides[1].left_items)
        self.assertEqual(["Автоматическая маршрутизация", "Сохранение памяти"], slides[1].right_items)
        self.assertEqual(["1 - единый интерфейс"], slides[1].metrics)
        self.assertEqual("Сотрудник работает в едином корпоративном AI-чате", slides[1].image_prompt)

    async def test_generate_slide_specs_enriches_image_slide_when_prompt_present(self) -> None:
        strategy = PresentationStrategy()
        fake_client = _FakePresentationClient(
            {
                "slides": [
                    {
                        "title": "GPTHub",
                        "subtitle": "Единая точка входа",
                        "bullets": ["Автоматическая маршрутизация"],
                    },
                    {
                        "title": "Рабочий сценарий",
                        "layout": "image_text",
                        "body": "Сотрудник задает вопрос и получает готовый результат в одном окне.",
                        "image_prompt": "Корпоративный сотрудник работает с AI-ассистентом в чате",
                    },
                ]
            }
        )

        with (
            patch("app.strategies.presentation.mws_client", fake_client),
            patch.object(strategy, "_fetch_image_bytes", new=AsyncMock(return_value=b"image-bytes")),
        ):
            slides = await strategy._generate_slide_specs(
                _strategy_request(task_type=TaskType.PRESENTATION, text="Сделай презентацию про GPTHub")
            )

        self.assertEqual(2, len(slides))
        self.assertEqual("https://cdn.example/presentation.png", slides[1].image_url)
        self.assertEqual(b"image-bytes", slides[1].image_bytes)
        self.assertEqual(settings.image_generation_model, fake_client.image_calls[0]["model"])
        self.assertIn("no text", fake_client.image_calls[0]["prompt"])
        self.assertIn("spot illustration", fake_client.image_calls[0]["prompt"])
        self.assertIn("not a full-slide background", fake_client.image_calls[0]["prompt"])

    async def test_enrich_slides_converts_extra_image_layouts_to_content(self) -> None:
        strategy = PresentationStrategy()
        strategy.max_generated_images = 1
        slides = [
            SlideSpec(
                title="Материалы",
                bullets=[],
                body="Дети рассматривают лейку, совок, землю и рассаду перед началом занятия.",
                image_prompt="Плоская иконка лейки, совка и рассады",
                layout="image_text",
            ),
            SlideSpec(
                title="Полив после посадки",
                bullets=["Льем воду ближе к корню", "Не размываем землю", "Проверяем влажность"],
                takeaway="Аккуратный полив помогает цветку прижиться.",
                image_prompt="Плоская иконка лейки рядом с цветком",
                layout="image_text",
            ),
        ]

        with (
            patch.object(strategy, "_generate_image_url", new=AsyncMock(return_value="https://cdn.example/icon.png")) as generate,
            patch.object(strategy, "_fetch_image_bytes", new=AsyncMock(return_value=b"image-bytes")),
        ):
            enriched = await strategy._enrich_slides_with_images(slides)

        self.assertEqual("image_text", enriched[0].layout)
        self.assertEqual(b"image-bytes", enriched[0].image_bytes)
        self.assertEqual("content", enriched[1].layout)
        self.assertIsNone(enriched[1].image_bytes)
        self.assertEqual("", enriched[1].image_url)
        self.assertEqual(1, generate.call_count)

    def test_build_pptx_uses_designed_widescreen_layout(self) -> None:
        strategy = PresentationStrategy()

        deck_bytes = strategy._build_pptx(
            [
                SlideSpec(
                    title="Архитектура GPTHub",
                    subtitle="Поток запроса от UI до модели",
                    bullets=["Frontend", "Router", "Strategy"],
                    takeaway="Маршрутизация отделяет UX от выбора модели.",
                    visual_hint="Схема из трех блоков",
                    layout="title",
                ),
                SlideSpec(
                    title="Ключевой сценарий",
                    bullets=["Классификация запроса", "Выбор стратегии", "Возврат результата"],
                    takeaway="Каждый шаг явно отделен и тестируем.",
                    visual_hint="Карточки по этапам",
                    layout="process",
                ),
                SlideSpec(
                    title="Пользовательский эффект",
                    bullets=[],
                    body="Пользователь получает результат в одном окне, а система сама подбирает нужную стратегию и модель.",
                    takeaway="Один сценарий заменяет ручное переключение между инструментами.",
                    layout="text",
                ),
                SlideSpec(
                    title="Визуальный пример",
                    bullets=["Единый чат", "Автоматический выбор"],
                    body="Иллюстрация помогает показать идею без дополнительного текста.",
                    image_prompt="Единый корпоративный чат с AI-ассистентом",
                    layout="image_text",
                ),
            ]
        )

        deck = Presentation(BytesIO(deck_bytes))

        self.assertEqual(4, len(deck.slides))
        self.assertGreater(deck.slide_width, deck.slide_height)
        self.assertGreater(len(deck.slides[0].shapes), 8)
        self.assertGreater(len(deck.slides[1].shapes), 12)
        self.assertGreater(len(deck.slides[2].shapes), 6)
        self.assertGreater(len(deck.slides[3].shapes), 6)
        visible_text = "\n".join(
            shape.text
            for slide in deck.slides
            for shape in slide.shapes
            if hasattr(shape, "text")
        )
        self.assertNotIn("Визуальный слайд", visible_text)
        self.assertNotIn("Единый корпоративный чат с AI-ассистентом", visible_text)
        self.assertNotIn("Схема из трех блоков", visible_text)


class ResponseMetadataTests(unittest.TestCase):
    def test_stream_chunk_includes_file_metadata_without_task_id(self) -> None:
        response = StrategyResponse(
            content="Презентация готова",
            model_used="slides-model",
            task_type=TaskType.PRESENTATION,
            routing_reason="done",
            file_url="/v1/files/presentation-file-id",
            sources=["/v1/files/presentation-file-id"],
        )

        payload = json.loads(
            stream_chunk(
                response,
                response.content,
                finish_reason=None,
                include_gpthub=True,
            )
            .decode("utf-8")
            .removeprefix("data: ")
        )

        self.assertEqual("/v1/files/presentation-file-id", payload["gpthub"]["file_url"])
        self.assertEqual(["/v1/files/presentation-file-id"], payload["gpthub"]["sources"])


class FileDownloadRouteTests(unittest.TestCase):
    def test_upload_returns_429_when_file_quota_is_exhausted(self) -> None:
        app = FastAPI()
        app.include_router(files_module.router, prefix="/v1")

        class _QuotaStorage:
            async def upload(self, **kwargs):
                raise FileQuotaExceededError("quota reached")

        with patch.object(files_module, "file_storage", _QuotaStorage()):
            response = TestClient(app).post(
                "/v1/files",
                headers={"x-user-id": "user-1"},
                files={"file": ("deck.pptx", b"x", "application/octet-stream")},
            )

        self.assertEqual(429, response.status_code)
        self.assertEqual("quota reached", response.json()["detail"])

    def test_download_uses_anonymous_user_when_header_is_missing(self) -> None:
        app = FastAPI()
        app.include_router(files_module.router, prefix="/v1")
        fake_storage = _FakeFileStorage()

        with patch.object(files_module, "file_storage", fake_storage):
            response = TestClient(app).get("/v1/files/presentation-file-id")

        self.assertEqual(200, response.status_code)
        self.assertEqual(b"pptx-bytes", response.content)
        self.assertEqual(
            [{"file_id": "presentation-file-id", "user_id": "anonymous"}],
            fake_storage.download_calls,
        )

    def test_download_with_access_token_bypasses_user_header(self) -> None:
        app = FastAPI()
        app.include_router(files_module.router, prefix="/v1")
        fake_storage = _FakeFileStorage()
        access_token = build_file_access_token("presentation-file-id")

        with patch.object(files_module, "file_storage", fake_storage):
            response = TestClient(app).get(f"/v1/files/presentation-file-id?access_token={access_token}")

        self.assertEqual(200, response.status_code)
        self.assertEqual(b"pptx-bytes", response.content)
        self.assertEqual([{"file_id": "presentation-file-id"}], fake_storage.shared_download_calls)
        self.assertEqual([], fake_storage.download_calls)


if __name__ == "__main__":
    unittest.main()
