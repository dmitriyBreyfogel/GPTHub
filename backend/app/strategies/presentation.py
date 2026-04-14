from __future__ import annotations

import base64
from dataclasses import dataclass, field, replace
from io import BytesIO
from typing import AsyncIterator

import httpx
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from app.core.config import settings
from app.core.prompt_cache import prompt_cache_manager
from app.providers.mws_gpt import ChatMessage, mws_client
from app.storage.files import build_file_access_token, file_storage
from app.strategies.base import StrategyRequest, StrategyResponse, TaskType
from app.strategies.response_utils import extract_json_object, stream_chunk


@dataclass(frozen=True)
class SlideSpec:
    title: str
    bullets: list[str]
    speaker_notes: str = ""
    subtitle: str = ""
    takeaway: str = ""
    visual_hint: str = ""
    layout: str = "content"
    body: str = ""
    left_title: str = ""
    right_title: str = ""
    left_items: list[str] = field(default_factory=list)
    right_items: list[str] = field(default_factory=list)
    metrics: list[str] = field(default_factory=list)
    image_prompt: str = ""
    image_url: str = ""
    image_bytes: bytes | None = field(default=None, repr=False, compare=False)


class PresentationStrategy:
    task_type = TaskType.PRESENTATION
    max_slides = 12
    max_generated_images = 2
    max_image_bytes = 8 * 1024 * 1024

    _slide_width = 13.333
    _slide_height = 7.5
    _supported_layouts = {
        "title",
        "section",
        "content",
        "two_column",
        "comparison",
        "timeline",
        "metrics",
        "statement",
        "text",
        "image",
        "image_text",
        "quote",
        "process",
    }
    _layout_aliases = {
        "image_with_text": "image_text",
        "text_image": "image_text",
        "picture": "image",
        "photo": "image",
        "big_statement": "statement",
        "key_statement": "statement",
        "plain_text": "text",
        "steps": "process",
    }
    _accent_colors = [
        RGBColor(227, 6, 19),
        RGBColor(0, 132, 112),
        RGBColor(0, 102, 204),
        RGBColor(112, 118, 128),
        RGBColor(118, 73, 160),
    ]
    _dark = RGBColor(18, 24, 32)
    _ink = RGBColor(24, 31, 42)
    _muted = RGBColor(91, 103, 119)
    _paper = RGBColor(247, 248, 251)
    _card = RGBColor(255, 255, 255)
    _line = RGBColor(224, 229, 237)
    _white = RGBColor(255, 255, 255)

    async def execute(self, request: StrategyRequest) -> StrategyResponse:
        if not request.text.strip():
            raise ValueError("Presentation topic is required")

        slides = await self._generate_slide_specs(request)
        pptx_bytes = self._build_pptx(slides)
        filename = self._filename(request.text)
        stored = await file_storage.upload(
            file_bytes=pptx_bytes,
            filename=filename,
            content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            user_id=request.user_id,
        )
        file_url = f"/v1/files/{stored.file_id}?access_token={build_file_access_token(stored.file_id)}"
        image_count = sum(1 for slide in slides if slide.image_bytes or slide.image_url)
        content = (
            "Презентация готова.\n"
            f"Файл: {filename}\n"
            f"[Скачать презентацию]({file_url})\n"
            f"Слайдов: {len(slides)}"
        )
        if image_count:
            content += f"\nИзображений: {image_count}"
        return StrategyResponse(
            content=content,
            model_used=request.model_override or settings.default_text_model,
            task_type=self.task_type,
            routing_reason="Presentation strategy: структура слайдов сгенерирована LLM, изображения добавлены по тематическим промптам, .pptx собран через python-pptx и сохранен в file storage.",
            file_url=file_url,
            sources=[file_url],
        )

    async def stream(self, request: StrategyRequest) -> AsyncIterator[bytes]:
        response = await self.execute(request)
        yield stream_chunk(response, response.content, finish_reason=None, include_gpthub=True)
        yield stream_chunk(response, "", finish_reason="stop", include_gpthub=True)
        yield b"data: [DONE]\n\n"

    async def _generate_slide_specs(self, request: StrategyRequest) -> list[SlideSpec]:
        messages = [
            ChatMessage(
                role="system",
                content=prompt_cache_manager.build_presentation_system_prompt(
                    workspace_instructions=request.workspace_instructions,
                ),
            ),
            ChatMessage(role="user", content=request.text.strip()),
        ]
        response = await mws_client.chat(
            messages,
            model=request.model_override,
            temperature=0.45,
            generation_options=request.generation_options,
        )
        data = extract_json_object(response.content) or {}
        slides = self._parse_slides(data)
        if not slides:
            slides = self._fallback_slides(request.text.strip())
        slides = slides[: self.max_slides]
        return await self._enrich_slides_with_images(slides)

    def _parse_slides(self, data: dict) -> list[SlideSpec]:
        raw_slides = data.get("slides")
        if not isinstance(raw_slides, list):
            return []

        slides = []
        for index, raw_slide in enumerate(raw_slides):
            if not isinstance(raw_slide, dict):
                continue
            title = self._clean_text(raw_slide.get("title"))
            bullets = self._parse_text_list(raw_slide.get("bullets"))
            subtitle = self._clean_text(raw_slide.get("subtitle"))
            takeaway = self._clean_text(raw_slide.get("takeaway"))
            visual_hint = self._clean_text(raw_slide.get("visual_hint"))
            speaker_notes = self._clean_text(raw_slide.get("speaker_notes"))
            layout = self._parse_layout(raw_slide.get("layout"), index=index)
            body = self._clean_body(raw_slide.get("body"))
            left_title = self._clean_text(raw_slide.get("left_title"))
            right_title = self._clean_text(raw_slide.get("right_title"))
            left_items = self._parse_text_list(raw_slide.get("left_items"))
            right_items = self._parse_text_list(raw_slide.get("right_items"))
            metrics = self._parse_metrics(raw_slide.get("metrics"))
            image_prompt = self._clean_text(raw_slide.get("image_prompt"))
            image_url = self._clean_text(raw_slide.get("image_url"))
            if title or bullets or body or takeaway:
                slides.append(
                    SlideSpec(
                        title=title or "Слайд",
                        bullets=bullets[:6],
                        speaker_notes=speaker_notes,
                        subtitle=subtitle,
                        takeaway=takeaway,
                        visual_hint=visual_hint,
                        layout=layout,
                        body=body,
                        left_title=left_title,
                        right_title=right_title,
                        left_items=left_items[:6],
                        right_items=right_items[:6],
                        metrics=metrics[:6],
                        image_prompt=image_prompt,
                        image_url=image_url,
                    )
                )
        return slides

    def _parse_text_list(self, raw_items: object) -> list[str]:
        if isinstance(raw_items, list):
            return [item for item in (self._clean_text(item) for item in raw_items) if item]
        if isinstance(raw_items, str):
            return [line.strip("-• \t") for line in raw_items.splitlines() if line.strip("-• \t")]
        return []

    def _parse_metrics(self, raw_items: object) -> list[str]:
        if not isinstance(raw_items, list):
            return self._parse_text_list(raw_items)

        metrics = []
        for item in raw_items:
            if isinstance(item, dict):
                value = self._clean_text(item.get("value"))
                label = self._clean_text(item.get("label"))
                note = self._clean_text(item.get("note"))
                metric = " - ".join(part for part in (value, label, note) if part)
            else:
                metric = self._clean_text(item)
            if metric:
                metrics.append(metric)
        return metrics

    def _parse_layout(self, raw_layout: object, *, index: int) -> str:
        layout = self._clean_text(raw_layout).lower().replace("-", "_").replace(" ", "_")
        layout = self._layout_aliases.get(layout, layout)
        if index == 0:
            return "title"
        if layout in self._supported_layouts:
            return layout
        return "content"

    async def _enrich_slides_with_images(self, slides: list[SlideSpec]) -> list[SlideSpec]:
        enriched = []
        generated = 0
        for slide in slides:
            if not self._should_generate_image(slide):
                enriched.append(slide)
                continue
            if generated >= self.max_generated_images:
                enriched.append(self._content_layout_for_image_fallback(slide))
                continue

            prompt = self._image_generation_prompt(slide)
            try:
                image_url = slide.image_url or await self._generate_image_url(prompt)
                image_bytes = await self._fetch_image_bytes(image_url)
            except Exception:
                enriched.append(self._content_layout_for_image_fallback(slide))
                continue

            if image_bytes:
                generated += 1
                enriched.append(replace(slide, image_url=image_url, image_bytes=image_bytes))
            else:
                enriched.append(self._content_layout_for_image_fallback(slide))
        return enriched

    def _should_generate_image(self, slide: SlideSpec) -> bool:
        return slide.layout in {"image", "image_text"} and bool(slide.image_prompt or slide.visual_hint)

    def _content_layout_for_image_fallback(self, slide: SlideSpec) -> SlideSpec:
        if slide.layout not in {"image", "image_text"}:
            return slide

        text_content = slide.body or slide.takeaway or slide.subtitle
        layout = "text" if text_content and not slide.bullets else "content"
        body = slide.body or (text_content if layout == "text" else "")
        return replace(slide, layout=layout, body=body, image_url="", image_bytes=None)

    async def _generate_image_url(self, prompt: str) -> str:
        try:
            return await mws_client.generate_image(prompt, model=settings.image_generation_model)
        except Exception:
            fallback_model = settings.image_generation_fallback_model
            if not fallback_model or fallback_model.lower() == settings.image_generation_model.lower():
                raise
            return await mws_client.generate_image(prompt, model=fallback_model)

    async def _fetch_image_bytes(self, image_url: str) -> bytes | None:
        if not image_url:
            return None
        if image_url.startswith("data:image/") and ";base64," in image_url:
            _, encoded = image_url.split(";base64,", 1)
            decoded = base64.b64decode(encoded)
            if len(decoded) > self.max_image_bytes:
                return None
            return decoded
        if not image_url.startswith(("http://", "https://")):
            return None

        async with httpx.AsyncClient(timeout=30.0, follow_redirects=True) as client:
            response = await client.get(image_url)
            response.raise_for_status()
            content_type = response.headers.get("content-type", "")
            if "image/" not in content_type.lower():
                return None
            if len(response.content) > self.max_image_bytes:
                return None
            return response.content

    def _image_generation_prompt(self, slide: SlideSpec) -> str:
        prompt = slide.image_prompt or slide.visual_hint or slide.title
        return (
            f"{prompt}. "
            "Simple flat vector spot illustration or icon for a presentation slide, isolated on a light background, "
            "friendly minimal style, no text, no letters, no logos, no watermark, not photorealistic, "
            "not a full-slide background."
        )

    def _fallback_slides(self, topic: str) -> list[SlideSpec]:
        normalized_topic = topic[:80] or "Презентация"
        return [
            SlideSpec(
                title=normalized_topic,
                bullets=["Контекст", "Цель", "Ожидаемый результат"],
                subtitle="Короткая, структурная версия для обсуждения",
                takeaway="Презентация задает общий контекст и фокус дальнейшего разговора.",
                visual_hint="Обложка с крупным заголовком и акцентной полосой",
                layout="title",
            ),
            SlideSpec(
                title="Главная мысль",
                bullets=[],
                takeaway="Одна сильная идея помогает аудитории быстро понять, зачем нужен разговор.",
                visual_hint="Крупная типографика на темном фоне",
                layout="statement",
            ),
            SlideSpec(
                title="Проблема и решение",
                bullets=["Текущая ситуация", "Ключевые ограничения", "Что меняется после решения"],
                left_title="Сейчас",
                right_title="После",
                left_items=["Разрозненный контекст", "Нечеткий сценарий", "Много ручных действий"],
                right_items=["Единая логика", "Понятный пользовательский путь", "Проверяемый результат"],
                takeaway="Решение должно быть объяснимым через сценарий и измеримую ценность.",
                layout="comparison",
            ),
            SlideSpec(
                title="Как это работает",
                bullets=["Сформировать запрос", "Выбрать подход", "Собрать результат"],
                takeaway="Процесс должен быть простым для пользователя и управляемым для команды.",
                visual_hint="Таймлайн из трех этапов",
                layout="process",
            ),
            SlideSpec(
                title="Следующие шаги",
                body="Сначала стоит проверить главный риск, затем собрать короткое демо и только после этого расширять функциональность. Такой порядок позволяет быстро показать ценность и не распыляться на второстепенные детали.",
                bullets=[],
                takeaway="Следующий шаг должен сокращать главный риск и приближать демонстрацию ценности.",
                layout="text",
            ),
        ]

    def _build_pptx(self, slides: list[SlideSpec]) -> bytes:
        if not slides:
            slides = self._fallback_slides("Презентация")

        deck = Presentation()
        deck.slide_width = Inches(self._slide_width)
        deck.slide_height = Inches(self._slide_height)
        total = len(slides)

        self._add_title_slide(deck, slides[0], slide_number=1, total_slides=total)
        for index, slide_spec in enumerate(slides[1:], start=2):
            self._add_slide_by_layout(deck, slide_spec, slide_number=index, total_slides=total)

        output = BytesIO()
        deck.save(output)
        return output.getvalue()

    def _add_slide_by_layout(
        self,
        deck: Presentation,
        spec: SlideSpec,
        *,
        slide_number: int,
        total_slides: int,
    ) -> None:
        if spec.layout == "section":
            self._add_section_slide(deck, spec, slide_number=slide_number, total_slides=total_slides)
        elif spec.layout in {"statement", "quote"}:
            self._add_statement_slide(deck, spec, slide_number=slide_number, total_slides=total_slides)
        elif spec.layout == "text":
            self._add_text_slide(deck, spec, slide_number=slide_number, total_slides=total_slides)
        elif spec.layout == "two_column":
            self._add_two_column_slide(deck, spec, slide_number=slide_number, total_slides=total_slides)
        elif spec.layout == "comparison":
            self._add_comparison_slide(deck, spec, slide_number=slide_number, total_slides=total_slides)
        elif spec.layout == "timeline":
            self._add_timeline_slide(deck, spec, slide_number=slide_number, total_slides=total_slides)
        elif spec.layout == "metrics":
            self._add_metrics_slide(deck, spec, slide_number=slide_number, total_slides=total_slides)
        elif spec.layout == "process":
            self._add_process_slide(deck, spec, slide_number=slide_number, total_slides=total_slides)
        elif spec.layout in {"image", "image_text"}:
            self._add_image_slide(deck, spec, slide_number=slide_number, total_slides=total_slides)
        else:
            self._add_content_slide(deck, spec, slide_number=slide_number, total_slides=total_slides)

    def _add_title_slide(
        self,
        deck: Presentation,
        spec: SlideSpec,
        *,
        slide_number: int,
        total_slides: int,
    ) -> None:
        slide = self._blank_slide(deck)
        accent = self._accent_for(slide_number)
        self._shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, self._slide_width, self._slide_height, self._dark)
        self._shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.28, self._slide_height, accent)
        self._shape(slide, MSO_SHAPE.RECTANGLE, 0.85, 0.8, 1.5, 0.08, accent)
        self._shape(slide, MSO_SHAPE.RECTANGLE, 9.9, 0, 3.45, self._slide_height, accent)

        self._textbox(
            slide,
            spec.title,
            0.85,
            1.35,
            8.4,
            1.45,
            font_size=42,
            color=self._white,
            bold=True,
        )
        subtitle = spec.subtitle or self._derive_takeaway(spec)
        self._textbox(
            slide,
            subtitle,
            0.9,
            3.0,
            7.8,
            0.85,
            font_size=18,
            color=RGBColor(218, 226, 237),
        )

        for index, bullet in enumerate((spec.bullets or [])[:3]):
            self._pill(slide, bullet, 0.9, 4.35 + index * 0.62, 7.25, 0.42, RGBColor(38, 48, 63))

        insight = spec.takeaway or "Ключевая мысль презентации"
        self._textbox(slide, "Фокус", 10.3, 1.18, 2.45, 0.38, font_size=12, color=self._white, bold=True)
        self._textbox(slide, insight, 10.3, 1.65, 2.15, 1.55, font_size=19, color=self._white, bold=True)
        self._add_footer(slide, slide_number, total_slides, accent, dark=True)
        self._apply_notes(slide, spec.speaker_notes)

    def _add_section_slide(
        self,
        deck: Presentation,
        spec: SlideSpec,
        *,
        slide_number: int,
        total_slides: int,
    ) -> None:
        slide = self._blank_slide(deck)
        accent = self._accent_for(slide_number)
        self._shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, self._slide_width, self._slide_height, self._dark)
        self._shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, self._slide_width, 0.18, accent)
        self._textbox(slide, f"{slide_number:02d}", 0.9, 1.25, 1.5, 0.8, font_size=28, color=accent, bold=True)
        self._textbox(slide, spec.title, 0.9, 2.25, 9.8, 1.2, font_size=38, color=self._white, bold=True)
        if spec.subtitle or spec.takeaway:
            self._textbox(
                slide,
                spec.subtitle or spec.takeaway,
                0.95,
                3.65,
                8.8,
                0.95,
                font_size=18,
                color=RGBColor(218, 226, 237),
            )
        self._add_footer(slide, slide_number, total_slides, accent, dark=True)
        self._apply_notes(slide, spec.speaker_notes)

    def _add_content_slide(
        self,
        deck: Presentation,
        spec: SlideSpec,
        *,
        slide_number: int,
        total_slides: int,
    ) -> None:
        slide = self._blank_slide(deck)
        accent = self._accent_for(slide_number)
        self._shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, self._slide_width, self._slide_height, self._paper)
        self._shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.18, self._slide_height, accent)
        self._shape(slide, MSO_SHAPE.RECTANGLE, 0.78, 0.58, 1.2, 0.08, accent)

        self._textbox(slide, spec.title, 0.78, 0.72, 8.45, 0.55, font_size=26, color=self._ink, bold=True)
        if spec.subtitle:
            self._textbox(slide, spec.subtitle, 0.8, 1.25, 8.1, 0.36, font_size=12, color=self._muted)

        self._add_bullet_cards(slide, spec, accent)
        self._add_insight_panel(slide, spec, accent)
        self._add_footer(slide, slide_number, total_slides, accent, dark=False)
        self._apply_notes(slide, spec.speaker_notes)

    def _add_statement_slide(
        self,
        deck: Presentation,
        spec: SlideSpec,
        *,
        slide_number: int,
        total_slides: int,
    ) -> None:
        slide = self._blank_slide(deck)
        accent = self._accent_for(slide_number)
        self._shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, self._slide_width, self._slide_height, self._dark)
        self._shape(slide, MSO_SHAPE.RECTANGLE, 0.78, 0.78, 0.1, 5.8, accent)
        self._textbox(slide, spec.title, 1.05, 0.85, 6.9, 0.45, font_size=15, color=accent, bold=True)
        statement = spec.takeaway or spec.body or self._derive_takeaway(spec)
        self._textbox(slide, statement, 1.0, 2.05, 10.75, 1.8, font_size=34, color=self._white, bold=True)
        if spec.body and spec.body != statement:
            self._textbox(slide, spec.body, 1.08, 4.55, 9.6, 0.9, font_size=16, color=RGBColor(218, 226, 237))
        elif spec.subtitle:
            self._textbox(slide, spec.subtitle, 1.08, 4.55, 9.6, 0.9, font_size=16, color=RGBColor(218, 226, 237))
        self._shape(slide, MSO_SHAPE.OVAL, 10.7, 0.6, 2.55, 2.55, RGBColor(42, 53, 70))
        self._shape(slide, MSO_SHAPE.OVAL, 11.15, 1.05, 1.65, 1.65, accent)
        self._add_footer(slide, slide_number, total_slides, accent, dark=True)
        self._apply_notes(slide, spec.speaker_notes)

    def _add_text_slide(
        self,
        deck: Presentation,
        spec: SlideSpec,
        *,
        slide_number: int,
        total_slides: int,
    ) -> None:
        slide = self._blank_slide(deck)
        accent = self._accent_for(slide_number)
        self._shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, self._slide_width, self._slide_height, self._paper)
        self._shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.18, self._slide_height, accent)
        self._textbox(slide, spec.title, 0.9, 0.78, 6.9, 0.58, font_size=28, color=self._ink, bold=True)
        if spec.subtitle:
            self._textbox(slide, spec.subtitle, 0.92, 1.35, 7.1, 0.38, font_size=13, color=self._muted)
        body = spec.body or " ".join(spec.bullets) or spec.takeaway
        self._shape(slide, MSO_SHAPE.RECTANGLE, 0.92, 2.0, 0.08, 3.82, accent)
        self._textbox(slide, body, 1.18, 2.02, 7.4, 3.65, font_size=22, color=self._ink, bold=True)
        self._shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 9.35, 1.28, 2.95, 4.85, self._dark)
        self._textbox(slide, "Вывод", 9.72, 1.72, 2.2, 0.32, font_size=12, color=accent, bold=True)
        self._textbox(slide, spec.takeaway or self._derive_takeaway(spec), 9.72, 2.25, 2.18, 1.5, font_size=18, color=self._white, bold=True)
        self._add_footer(slide, slide_number, total_slides, accent, dark=False)
        self._apply_notes(slide, spec.speaker_notes)

    def _add_two_column_slide(
        self,
        deck: Presentation,
        spec: SlideSpec,
        *,
        slide_number: int,
        total_slides: int,
    ) -> None:
        slide = self._blank_slide(deck)
        accent = self._accent_for(slide_number)
        self._add_light_background(slide, accent)
        self._add_header(slide, spec, accent)
        left_items, right_items = self._column_items(spec)
        self._column_panel(slide, spec.left_title or "Первый блок", left_items, 0.82, 1.85, 5.62, 4.45, accent)
        self._column_panel(slide, spec.right_title or "Второй блок", right_items, 6.72, 1.85, 5.62, 4.45, self._accent_for(slide_number + 1))
        self._takeaway_band(slide, spec, accent)
        self._add_footer(slide, slide_number, total_slides, accent, dark=False)
        self._apply_notes(slide, spec.speaker_notes)

    def _add_comparison_slide(
        self,
        deck: Presentation,
        spec: SlideSpec,
        *,
        slide_number: int,
        total_slides: int,
    ) -> None:
        slide = self._blank_slide(deck)
        accent = self._accent_for(slide_number)
        secondary = self._accent_for(slide_number + 2)
        self._add_light_background(slide, accent)
        self._add_header(slide, spec, accent)
        left_items, right_items = self._column_items(spec)
        self._comparison_panel(slide, spec.left_title or "Было", left_items, 0.82, 1.9, 5.55, 4.35, accent)
        self._comparison_panel(slide, spec.right_title or "Стало", right_items, 6.85, 1.9, 5.55, 4.35, secondary)
        self._textbox(slide, "vs", 6.23, 3.62, 0.65, 0.35, font_size=18, color=self._muted, bold=True, align=PP_ALIGN.CENTER)
        self._takeaway_band(slide, spec, accent)
        self._add_footer(slide, slide_number, total_slides, accent, dark=False)
        self._apply_notes(slide, spec.speaker_notes)

    def _add_timeline_slide(
        self,
        deck: Presentation,
        spec: SlideSpec,
        *,
        slide_number: int,
        total_slides: int,
    ) -> None:
        slide = self._blank_slide(deck)
        accent = self._accent_for(slide_number)
        self._add_light_background(slide, accent)
        self._add_header(slide, spec, accent)
        steps = (spec.left_items or spec.bullets or [spec.takeaway or spec.title])[:5]
        self._shape(slide, MSO_SHAPE.RECTANGLE, 1.16, 3.18, 10.9, 0.08, self._line)
        for index, step in enumerate(steps):
            left = 1.05 + index * (9.85 / max(len(steps) - 1, 1))
            self._shape(slide, MSO_SHAPE.OVAL, left, 2.85, 0.72, 0.72, accent)
            self._textbox(slide, str(index + 1), left, 3.05, 0.72, 0.2, font_size=10, color=self._white, bold=True, align=PP_ALIGN.CENTER)
            self._textbox(slide, step, max(0.75, left - 0.55), 3.85, 1.88, 0.92, font_size=13, color=self._ink, bold=True, align=PP_ALIGN.CENTER)
        self._takeaway_band(slide, spec, accent)
        self._add_footer(slide, slide_number, total_slides, accent, dark=False)
        self._apply_notes(slide, spec.speaker_notes)

    def _add_metrics_slide(
        self,
        deck: Presentation,
        spec: SlideSpec,
        *,
        slide_number: int,
        total_slides: int,
    ) -> None:
        slide = self._blank_slide(deck)
        accent = self._accent_for(slide_number)
        self._add_light_background(slide, accent)
        self._add_header(slide, spec, accent)
        metrics = (spec.metrics or spec.bullets or [spec.takeaway or spec.title])[:4]
        for index, metric in enumerate(metrics):
            value, label = self._split_metric(metric)
            left = 0.85 + (index % 2) * 6.0
            top = 1.9 + (index // 2) * 2.05
            color = self._accent_for(slide_number + index)
            self._shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, top, 5.45, 1.58, self._card, self._line)
            self._shape(slide, MSO_SHAPE.RECTANGLE, left, top, 0.12, 1.58, color)
            self._textbox(slide, value, left + 0.38, top + 0.34, 1.95, 0.58, font_size=25, color=color, bold=True)
            self._textbox(slide, label, left + 2.45, top + 0.38, 2.42, 0.72, font_size=14, color=self._ink, bold=True)
        self._takeaway_band(slide, spec, accent)
        self._add_footer(slide, slide_number, total_slides, accent, dark=False)
        self._apply_notes(slide, spec.speaker_notes)

    def _add_process_slide(
        self,
        deck: Presentation,
        spec: SlideSpec,
        *,
        slide_number: int,
        total_slides: int,
    ) -> None:
        slide = self._blank_slide(deck)
        accent = self._accent_for(slide_number)
        self._add_light_background(slide, accent)
        self._add_header(slide, spec, accent)
        steps = (spec.left_items or spec.bullets or [spec.takeaway or spec.title])[:5]
        card_width = 2.18 if len(steps) >= 5 else 2.55
        left_base = 0.82
        top = 2.26
        for index, step in enumerate(steps):
            left = left_base + index * (card_width + 0.22)
            color = self._accent_for(slide_number + index)
            self._shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_width, 2.05, self._card, self._line)
            self._shape(slide, MSO_SHAPE.OVAL, left + 0.28, top + 0.28, 0.58, 0.58, color)
            self._textbox(slide, str(index + 1), left + 0.28, top + 0.45, 0.58, 0.18, font_size=9, color=self._white, bold=True, align=PP_ALIGN.CENTER)
            self._textbox(slide, step, left + 0.26, top + 1.08, card_width - 0.52, 0.62, font_size=13, color=self._ink, bold=True, align=PP_ALIGN.CENTER)
            if index < len(steps) - 1:
                self._textbox(slide, "→", left + card_width - 0.02, top + 0.77, 0.36, 0.3, font_size=16, color=self._muted, bold=True)
        self._takeaway_band(slide, spec, accent)
        self._add_footer(slide, slide_number, total_slides, accent, dark=False)
        self._apply_notes(slide, spec.speaker_notes)

    def _add_image_slide(
        self,
        deck: Presentation,
        spec: SlideSpec,
        *,
        slide_number: int,
        total_slides: int,
    ) -> None:
        slide = self._blank_slide(deck)
        accent = self._accent_for(slide_number)
        self._add_light_background(slide, accent)
        self._add_header(slide, spec, accent)

        if spec.layout == "image_text":
            text = spec.body or spec.takeaway or self._derive_takeaway(spec)
            self._shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.82, 1.82, 7.12, 3.92, self._card, self._line)
            self._shape(slide, MSO_SHAPE.RECTANGLE, 0.82, 1.82, 0.14, 3.92, accent)
            self._textbox(slide, text, 1.18, 2.16, 6.25, 1.42, font_size=21, color=self._ink, bold=True)
            for index, bullet in enumerate((spec.bullets or [])[:3]):
                top = 4.02 + index * 0.5
                self._shape(slide, MSO_SHAPE.OVAL, 1.18, top + 0.05, 0.26, 0.26, accent)
                self._textbox(slide, bullet, 1.62, top, 5.75, 0.34, font_size=13, color=self._ink, bold=True)
            self._image_box(slide, spec, 8.48, 1.84, 3.72, 3.88, accent)
        else:
            self._image_box(slide, spec, 3.42, 1.72, 6.48, 3.72, accent)
            caption = spec.body or spec.takeaway or spec.subtitle
            if caption:
                self._textbox(slide, caption, 2.1, 5.75, 9.05, 0.38, font_size=15, color=self._ink, bold=True, align=PP_ALIGN.CENTER)

        self._takeaway_band(slide, spec, accent)
        self._add_footer(slide, slide_number, total_slides, accent, dark=False)
        self._apply_notes(slide, spec.speaker_notes)

    def _add_bullet_cards(self, slide, spec: SlideSpec, accent: RGBColor) -> None:
        bullets = (spec.bullets or [spec.takeaway or spec.title])[:6]
        columns = 2 if len(bullets) > 3 else 1
        card_width = 3.82 if columns == 2 else 7.95
        card_height = 1.02 if columns == 2 else 0.82
        left_base = 0.78
        top_base = 1.95
        horizontal_gap = 0.32
        vertical_gap = 0.24

        for index, bullet in enumerate(bullets):
            column = index % columns
            row = index // columns
            left = left_base + column * (card_width + horizontal_gap)
            top = top_base + row * (card_height + vertical_gap)
            self._shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_width, card_height, self._card, self._line)
            self._shape(slide, MSO_SHAPE.OVAL, left + 0.22, top + 0.26, 0.42, 0.42, accent)
            self._textbox(
                slide,
                str(index + 1),
                left + 0.22,
                top + 0.31,
                0.42,
                0.23,
                font_size=10,
                color=self._white,
                bold=True,
                align=PP_ALIGN.CENTER,
            )
            self._textbox(
                slide,
                bullet,
                left + 0.78,
                top + 0.22,
                card_width - 1.02,
                card_height - 0.36,
                font_size=14 if columns == 2 else 15,
                color=self._ink,
                bold=True,
            )

    def _add_insight_panel(self, slide, spec: SlideSpec, accent: RGBColor) -> None:
        self._shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 9.2, 1.05, 3.32, 5.3, self._dark)
        self._shape(slide, MSO_SHAPE.RECTANGLE, 9.2, 1.05, 0.14, 5.3, accent)
        self._textbox(slide, "Ключевая мысль", 9.55, 1.42, 2.45, 0.32, font_size=12, color=accent, bold=True)
        self._textbox(
            slide,
            spec.takeaway or self._derive_takeaway(spec),
            9.55,
            1.92,
            2.38,
            1.65,
            font_size=18,
            color=self._white,
            bold=True,
        )
        detail = spec.subtitle or spec.body or self._derive_takeaway(spec)
        self._textbox(slide, "Контекст", 9.55, 4.35, 2.45, 0.32, font_size=11, color=RGBColor(205, 214, 226), bold=True)
        self._textbox(slide, detail, 9.55, 4.78, 2.32, 0.85, font_size=12, color=RGBColor(224, 230, 239))

    def _add_light_background(self, slide, accent: RGBColor) -> None:
        self._shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, self._slide_width, self._slide_height, self._paper)
        self._shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.18, self._slide_height, accent)
        self._shape(slide, MSO_SHAPE.OVAL, 11.32, -0.72, 2.4, 2.4, RGBColor(236, 239, 245), RGBColor(236, 239, 245))

    def _add_header(self, slide, spec: SlideSpec, accent: RGBColor) -> None:
        self._shape(slide, MSO_SHAPE.RECTANGLE, 0.78, 0.58, 1.2, 0.08, accent)
        self._textbox(slide, spec.title, 0.78, 0.72, 8.45, 0.55, font_size=26, color=self._ink, bold=True)
        if spec.subtitle:
            self._textbox(slide, spec.subtitle, 0.8, 1.25, 8.1, 0.36, font_size=12, color=self._muted)

    def _column_panel(
        self,
        slide,
        title: str,
        items: list[str],
        left: float,
        top: float,
        width: float,
        height: float,
        accent: RGBColor,
    ) -> None:
        self._shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height, self._card, self._line)
        self._shape(slide, MSO_SHAPE.RECTANGLE, left, top, width, 0.16, accent)
        self._textbox(slide, title, left + 0.38, top + 0.42, width - 0.76, 0.38, font_size=17, color=self._ink, bold=True)
        for index, item in enumerate(items[:4]):
            y = top + 1.12 + index * 0.72
            self._shape(slide, MSO_SHAPE.OVAL, left + 0.42, y + 0.08, 0.28, 0.28, accent)
            self._textbox(slide, item, left + 0.88, y, width - 1.15, 0.38, font_size=13, color=self._ink, bold=True)

    def _comparison_panel(
        self,
        slide,
        title: str,
        items: list[str],
        left: float,
        top: float,
        width: float,
        height: float,
        accent: RGBColor,
    ) -> None:
        self._shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height, self._card, self._line)
        self._shape(slide, MSO_SHAPE.RECTANGLE, left, top, width, 0.5, accent)
        self._textbox(slide, title, left + 0.32, top + 0.16, width - 0.64, 0.24, font_size=14, color=self._white, bold=True)
        for index, item in enumerate(items[:4]):
            y = top + 0.95 + index * 0.76
            self._textbox(slide, item, left + 0.45, y, width - 0.9, 0.42, font_size=14, color=self._ink, bold=True)

    def _takeaway_band(self, slide, spec: SlideSpec, accent: RGBColor) -> None:
        takeaway = spec.takeaway or self._derive_takeaway(spec)
        self._shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.82, 6.35, 11.52, 0.52, self._dark)
        self._shape(slide, MSO_SHAPE.RECTANGLE, 0.82, 6.35, 0.12, 0.52, accent)
        self._textbox(slide, takeaway, 1.08, 6.49, 10.55, 0.22, font_size=11, color=self._white, bold=True)

    def _image_box(
        self,
        slide,
        spec: SlideSpec,
        left: float,
        top: float,
        width: float,
        height: float,
        accent: RGBColor,
        *,
        full_bleed: bool = False,
    ) -> None:
        if spec.image_bytes:
            try:
                slide.shapes.add_picture(
                    BytesIO(spec.image_bytes),
                    Inches(left),
                    Inches(top),
                    width=Inches(width),
                    height=Inches(height),
                )
                return
            except Exception:
                pass

        fill = RGBColor(244, 247, 250) if full_bleed else self._card
        self._shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height, fill, self._line)
        self._shape(slide, MSO_SHAPE.RECTANGLE, left, top, width, 0.12, accent)

        center_x = left + width / 2
        center_y = top + height / 2
        soft = RGBColor(228, 234, 242)
        muted = RGBColor(112, 118, 128)

        self._shape(slide, MSO_SHAPE.OVAL, center_x - 0.95, center_y - 0.95, 1.9, 1.9, soft, soft)
        self._shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, center_x - 0.62, center_y - 0.32, 1.24, 0.9, accent, accent)
        self._shape(slide, MSO_SHAPE.OVAL, center_x - 0.38, center_y - 0.03, 0.28, 0.28, self._white, self._white)
        self._shape(slide, MSO_SHAPE.OVAL, center_x + 0.1, center_y - 0.03, 0.28, 0.28, self._white, self._white)
        self._shape(slide, MSO_SHAPE.RECTANGLE, center_x - 0.42, center_y + 0.78, 0.84, 0.1, muted, muted)

    def _add_footer(self, slide, slide_number: int, total_slides: int, accent: RGBColor, *, dark: bool) -> None:
        track = RGBColor(58, 68, 82) if dark else self._line
        text_color = RGBColor(214, 222, 233) if dark else self._muted
        self._shape(slide, MSO_SHAPE.RECTANGLE, 0.78, 7.05, 10.9, 0.035, track)
        self._shape(slide, MSO_SHAPE.RECTANGLE, 0.78, 7.05, 10.9 * slide_number / max(total_slides, 1), 0.035, accent)
        self._textbox(
            slide,
            f"{slide_number:02d}/{total_slides:02d}",
            11.9,
            6.86,
            0.8,
            0.22,
            font_size=9,
            color=text_color,
            align=PP_ALIGN.RIGHT,
        )

    def _blank_slide(self, deck: Presentation):
        return deck.slides.add_slide(deck.slide_layouts[6])

    def _accent_for(self, slide_number: int) -> RGBColor:
        return self._accent_colors[(slide_number - 1) % len(self._accent_colors)]

    def _column_items(self, spec: SlideSpec) -> tuple[list[str], list[str]]:
        if spec.left_items or spec.right_items:
            left = spec.left_items or spec.bullets[: max(1, len(spec.bullets) // 2)]
            right = spec.right_items or spec.bullets[max(1, len(spec.bullets) // 2) :]
            return left or [spec.takeaway or spec.title], right or [spec.takeaway or spec.title]
        bullets = spec.bullets or [spec.takeaway or spec.title]
        midpoint = max(1, (len(bullets) + 1) // 2)
        return bullets[:midpoint], bullets[midpoint:] or [spec.takeaway or spec.title]

    def _split_metric(self, metric: str) -> tuple[str, str]:
        for separator in (" - ", ": "):
            if separator in metric:
                value, label = metric.split(separator, 1)
                return value.strip(), label.strip()
        parts = metric.split(maxsplit=1)
        if len(parts) == 2 and any(char.isdigit() for char in parts[0]):
            return parts[0], parts[1]
        return metric, "Ключевой показатель"

    def _shape(
        self,
        slide,
        shape_type: MSO_SHAPE,
        left: float,
        top: float,
        width: float,
        height: float,
        fill: RGBColor,
        line: RGBColor | None = None,
    ):
        shape = slide.shapes.add_shape(
            shape_type,
            Inches(left),
            Inches(top),
            Inches(width),
            Inches(height),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.color.rgb = line or fill
        return shape

    def _textbox(
        self,
        slide,
        text: str,
        left: float,
        top: float,
        width: float,
        height: float,
        *,
        font_size: int,
        color: RGBColor,
        bold: bool = False,
        align: PP_ALIGN = PP_ALIGN.LEFT,
    ):
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        frame.margin_left = Inches(0.03)
        frame.margin_right = Inches(0.03)
        paragraph = frame.paragraphs[0]
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = text
        run.font.name = "Aptos"
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = color
        return box

    def _pill(self, slide, text: str, left: float, top: float, width: float, height: float, fill: RGBColor) -> None:
        self._shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height, fill)
        self._textbox(slide, text, left + 0.22, top + 0.08, width - 0.44, height - 0.12, font_size=12, color=self._white)

    def _derive_takeaway(self, spec: SlideSpec) -> str:
        if spec.takeaway:
            return spec.takeaway
        if spec.body:
            return spec.body
        if spec.bullets:
            return spec.bullets[0]
        return spec.title

    def _visual_hint_for(self, spec: SlideSpec) -> str:
        if spec.layout == "timeline":
            return "Таймлайн этапов с короткими подписями."
        if spec.layout == "comparison":
            return "Сравнение двух вариантов в соседних блоках."
        if spec.layout == "metrics":
            return "Крупные метрики или индикаторы результата."
        if spec.layout == "two_column":
            return "Две смысловые колонки: контекст и решение."
        if spec.layout in {"image", "image_text"}:
            return "Тематическая иллюстрация, усиливающая ключевую мысль."
        return "Акцентная панель с главным выводом и карточками тезисов."

    def _apply_notes(self, slide, speaker_notes: str) -> None:
        if not speaker_notes:
            return
        notes_frame = slide.notes_slide.notes_text_frame
        notes_frame.text = speaker_notes

    def _filename(self, prompt: str) -> str:
        slug = "".join(char if char.isalnum() else "-" for char in prompt.lower())
        slug = "-".join(part for part in slug.split("-") if part)
        if not slug:
            slug = "presentation"
        return f"{slug[:60]}.pptx"

    def _clean_body(self, value: object) -> str:
        if isinstance(value, list):
            return "\n".join(item for item in (self._clean_text(item) for item in value) if item)
        return self._clean_text(value)

    def _clean_text(self, value: object) -> str:
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return str(value)
        if not isinstance(value, str):
            return ""
        return " ".join(value.split()).strip()
