from __future__ import annotations

import math
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import AsyncIterator

from docx import Document
from pypdf import PdfReader

from app.core.prompt_cache import prompt_cache_manager
from app.core.resource_limits import (
    MAX_FILE_QA_CHUNKS,
    MAX_FILE_QA_PDF_PAGES,
    MAX_FILE_QA_SOURCE_CHARS,
)
from app.core.response_formatting import append_technical_formatting_guidance
from app.providers.mws_gpt import ChatMessage, mws_client
from app.strategies.base import StrategyRequest, StrategyResponse, TaskType


@dataclass(frozen=True)
class RankedChunk:
    index: int
    text: str
    score: float


class FileQAStrategy:
    task_type = TaskType.FILE_QA
    chunk_size = 1800
    chunk_overlap = 250
    top_k = 5
    max_source_chars = MAX_FILE_QA_SOURCE_CHARS
    max_chunks = MAX_FILE_QA_CHUNKS
    max_pdf_pages = MAX_FILE_QA_PDF_PAGES

    async def execute(self, request: StrategyRequest) -> StrategyResponse:
        ranked_chunks = await self._rank_chunks(request)
        messages = self._build_messages(request, ranked_chunks)
        response = await mws_client.chat(
            messages,
            model=request.model_override,
            generation_options=request.generation_options,
        )
        return StrategyResponse(
            content=response.content,
            model_used=response.model,
            task_type=self.task_type,
            routing_reason=f"File QA strategy: документ разобран, выбрано {len(ranked_chunks)} релевантных чанков по cosine similarity.",
            sources=[f"chunk:{chunk.index}:score:{chunk.score:.4f}" for chunk in ranked_chunks],
        )

    async def stream(self, request: StrategyRequest) -> AsyncIterator[bytes]:
        ranked_chunks = await self._rank_chunks(request)
        messages = self._build_messages(request, ranked_chunks)
        async for chunk in mws_client.chat_stream(
            messages,
            model=request.model_override,
            generation_options=request.generation_options,
        ):
            yield chunk

    async def _rank_chunks(self, request: StrategyRequest) -> list[RankedChunk]:
        text = self._extract_text(request)
        chunks = self._chunk_text(text)
        if not chunks:
            raise ValueError("Document text is empty")

        query = request.text.strip() or "Сделай краткое резюме документа"
        query_embedding = await mws_client.embed(query)
        ranked_chunks = []

        for index, chunk in enumerate(chunks, start=1):
            chunk_embedding = await mws_client.embed(chunk)
            ranked_chunks.append(
                RankedChunk(
                    index=index,
                    text=chunk,
                    score=self._cosine_similarity(query_embedding, chunk_embedding),
                )
            )

        ranked_chunks.sort(key=lambda item: item.score, reverse=True)
        return ranked_chunks[: self.top_k]

    def _extract_text(self, request: StrategyRequest) -> str:
        if not request.file_bytes:
            raise ValueError("File bytes are required for file QA strategy")

        content_type = (request.file_content_type or "").lower()
        extension = Path(request.file_name or "").suffix.lower()

        if content_type == "application/pdf" or extension == ".pdf":
            return self._extract_pdf(request.file_bytes)
        if (
            content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            or extension == ".docx"
        ):
            return self._extract_docx(request.file_bytes)
        if content_type.startswith("text/") or extension in {".txt", ".md", ".csv", ".json"}:
            return self._extract_plain_text(request.file_bytes)

        return self._extract_plain_text(request.file_bytes)

    def _extract_pdf(self, file_bytes: bytes) -> str:
        reader = PdfReader(BytesIO(file_bytes))
        parts = []
        total_chars = 0
        for page_index, page in enumerate(reader.pages):
            if page_index >= self.max_pdf_pages:
                break
            page_text = page.extract_text() or ""
            normalized = self._normalize_text(page_text)
            if not normalized:
                continue
            remaining = self.max_source_chars - total_chars
            if remaining <= 0:
                break
            clipped = normalized[:remaining].rstrip()
            if clipped:
                parts.append(clipped)
                total_chars += len(clipped)
        return "\n\n".join(parts).strip()

    def _extract_docx(self, file_bytes: bytes) -> str:
        document = Document(BytesIO(file_bytes))
        parts = []
        total_chars = 0
        for paragraph in document.paragraphs:
            text = self._normalize_text(paragraph.text)
            if text:
                total_chars = self._append_limited_text(parts, text, total_chars)
                if total_chars >= self.max_source_chars:
                    break
        for table in document.tables:
            if total_chars >= self.max_source_chars:
                break
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    cell_text = self._normalize_text(cell.text)
                    if cell_text:
                        cells.append(cell_text)
                if cells:
                    total_chars = self._append_limited_text(parts, " | ".join(cells), total_chars)
                    if total_chars >= self.max_source_chars:
                        break
        return "\n".join(parts).strip()

    def _extract_plain_text(self, file_bytes: bytes) -> str:
        for encoding in ("utf-8", "utf-8-sig", "cp1251"):
            try:
                return self._normalize_text(file_bytes.decode(encoding))[: self.max_source_chars].rstrip()
            except UnicodeDecodeError:
                continue
        return self._normalize_text(file_bytes.decode("utf-8", errors="ignore"))[: self.max_source_chars].rstrip()

    def _chunk_text(self, text: str) -> list[str]:
        normalized = self._normalize_text(text)
        if not normalized:
            return []
        chunks = []
        start = 0
        while start < len(normalized):
            end = min(start + self.chunk_size, len(normalized))
            split_at = self._split_position(normalized, start, end)
            chunk = normalized[start:split_at].strip()
            if chunk:
                chunks.append(chunk)
            if len(chunks) >= self.max_chunks:
                break
            if split_at >= len(normalized):
                break
            start = max(0, split_at - self.chunk_overlap)
        return chunks

    def _append_limited_text(self, parts: list[str], text: str, total_chars: int) -> int:
        remaining = self.max_source_chars - total_chars
        if remaining <= 0:
            return total_chars
        clipped = text[:remaining].rstrip()
        if not clipped:
            return total_chars
        parts.append(clipped)
        return total_chars + len(clipped)

    def _split_position(self, text: str, start: int, end: int) -> int:
        if end >= len(text):
            return len(text)
        window = text[start:end]
        candidates = [window.rfind(separator) for separator in ("\n\n", "\n", ". ", "! ", "? ")]
        split = max(candidates)
        if split < self.chunk_size // 2:
            return end
        return start + split + 1

    def _build_messages(self, request: StrategyRequest, ranked_chunks: list[RankedChunk]) -> list[ChatMessage]:
        query = request.text.strip() or "Сделай краткое резюме документа"
        file_name = request.file_name or "uploaded file"
        context = "\n\n".join(
            f"[chunk {chunk.index}, score {chunk.score:.4f}]\n{chunk.text}"
            for chunk in ranked_chunks
        )
        return [
            ChatMessage(
                role="system",
                content=append_technical_formatting_guidance(
                    prompt_cache_manager.build_file_qa_system_prompt(
                        workspace_instructions=request.workspace_instructions
                    )
                ),
            ),
            ChatMessage(
                role="user",
                content=(
                    f"Файл: {file_name}\n\n"
                    f"Вопрос пользователя:\n{query}\n\n"
                    f"Релевантные фрагменты документа:\n{context}"
                ),
            ),
        ]

    def _normalize_text(self, text: str) -> str:
        lines = [" ".join(line.split()) for line in text.splitlines()]
        return "\n".join(line for line in lines if line).strip()

    def _cosine_similarity(self, left: list[float], right: list[float]) -> float:
        if not left or not right or len(left) != len(right):
            return 0.0
        dot = 0.0
        left_norm = 0.0
        right_norm = 0.0
        for left_value, right_value in zip(left, right):
            dot += left_value * right_value
            left_norm += left_value * left_value
            right_norm += right_value * right_value
        denom = math.sqrt(left_norm) * math.sqrt(right_norm)
        if denom == 0.0:
            return 0.0
        return dot / denom
