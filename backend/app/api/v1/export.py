from __future__ import annotations

import json
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Literal
from xml.sax.saxutils import escape

from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import Response
from pydantic import BaseModel, Field

from app.core.rate_limit import RateLimitRule, rate_limiter, request_subject
from app.core.resource_limits import (
    MAX_EXPORT_MESSAGES,
    MAX_EXPORT_TITLE_CHARS,
    MAX_EXPORT_TOTAL_CHARS,
    normalize_single_line_text,
)

router = APIRouter()

_EXPORT_RATE_LIMIT = RateLimitRule(
    scope="export",
    limit=20,
    window_seconds=600,
    detail="Export quota exceeded. Please retry later.",
)


class ExportMessage(BaseModel):
    role: str = "user"
    content: Any = ""
    created_at: str | None = None


class ExportRequest(BaseModel):
    messages: list[ExportMessage] = Field(default_factory=list)
    format: Literal["docx", "pdf"] = "docx"
    title: str = "Диалог GPTHub"


@router.post("/export")
async def export_dialog(request: Request, body: ExportRequest):
    await rate_limiter.enforce(subject=request_subject(request, "anonymous"), rule=_EXPORT_RATE_LIMIT)
    if len(body.messages) > MAX_EXPORT_MESSAGES:
        raise HTTPException(status_code=413, detail="Too many messages in export")

    messages = [_export_message(message) for message in body.messages]
    if not messages:
        raise HTTPException(status_code=400, detail="messages must not be empty")

    total_chars = sum(len(message["content"]) for message in messages)
    if total_chars > MAX_EXPORT_TOTAL_CHARS:
        raise HTTPException(status_code=413, detail="Export payload is too large")

    title = normalize_single_line_text(_clean_text(body.title), MAX_EXPORT_TITLE_CHARS) or "Диалог GPTHub"
    if body.format == "docx":
        content = _build_docx(title, messages)
        return _file_response(
            content,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename=_filename(title, "docx"),
        )

    content = _build_pdf(title, messages)
    return _file_response(
        content,
        media_type="application/pdf",
        filename=_filename(title, "pdf"),
    )


def _export_message(message: ExportMessage) -> dict[str, str]:
    return {
        "role": _clean_text(message.role) or "user",
        "content": _content_to_text(message.content),
        "created_at": _clean_text(message.created_at or ""),
    }


def _content_to_text(content: Any) -> str:
    if isinstance(content, str):
        return _clean_text(content)
    if isinstance(content, list):
        parts = []
        for item in content:
            if isinstance(item, str):
                parts.append(item)
            elif isinstance(item, dict):
                text = item.get("text") or item.get("content")
                if isinstance(text, str):
                    parts.append(text)
        text = "\n".join(parts).strip()
        if text:
            return _clean_text(text)
    if isinstance(content, dict):
        text = content.get("text") or content.get("content")
        if isinstance(text, str):
            return _clean_text(text)
    return _clean_text(json.dumps(content, ensure_ascii=False))


def _build_docx(title: str, messages: list[dict[str, str]]) -> bytes:
    try:
        from docx import Document
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="python-docx is not installed") from exc

    document = Document()
    document.add_heading(title, 0)
    document.add_paragraph(_exported_at())

    for message in messages:
        role = _role_label(message["role"])
        created_at = message["created_at"]
        heading = role if not created_at else f"{role} · {created_at}"
        paragraph = document.add_paragraph()
        run = paragraph.add_run(heading)
        run.bold = True
        document.add_paragraph(message["content"] or " ")

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _build_pdf(title: str, messages: list[dict[str, str]]) -> bytes:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="reportlab is not installed") from exc

    output = BytesIO()
    font_name = _pdf_font_name()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "GPTHubTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=18,
        leading=22,
        spaceAfter=10,
    )
    meta_style = ParagraphStyle(
        "GPTHubMeta",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=9,
        leading=12,
        spaceAfter=12,
    )
    role_style = ParagraphStyle(
        "GPTHubRole",
        parent=styles["Heading4"],
        fontName=font_name,
        fontSize=11,
        leading=14,
        spaceBefore=8,
        spaceAfter=4,
    )
    body_style = ParagraphStyle(
        "GPTHubBody",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=10,
        leading=14,
        spaceAfter=8,
    )
    story = [
        Paragraph(_pdf_text(title), title_style),
        Paragraph(_pdf_text(_exported_at()), meta_style),
    ]

    for message in messages:
        role = _role_label(message["role"])
        created_at = message["created_at"]
        heading = role if not created_at else f"{role} · {created_at}"
        story.append(Paragraph(_pdf_text(heading), role_style))
        story.append(Paragraph(_pdf_text(message["content"] or " "), body_style))
        story.append(Spacer(1, 2 * mm))

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
    )
    document.build(story)
    return output.getvalue()


def _pdf_font_name() -> str:
    try:
        import reportlab
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError:
        return "Helvetica"

    candidates = [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
        Path(reportlab.__file__).resolve().parent / "fonts" / "Vera.ttf",
    ]
    for path in candidates:
        if path.exists():
            font_name = f"GPTHubExportFont{abs(hash(str(path))) % 100000}"
            if font_name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(font_name, str(path)))
            return font_name
    return "Helvetica"


def _file_response(content: bytes, *, media_type: str, filename: str) -> Response:
    return Response(
        content=content,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _role_label(role: str) -> str:
    labels = {
        "system": "System",
        "user": "User",
        "assistant": "Assistant",
        "tool": "Tool",
    }
    return labels.get(role.strip().lower(), role.strip() or "Message")


def _filename(title: str, extension: str) -> str:
    slug = "".join(char.lower() if char.isascii() and char.isalnum() else "-" for char in title)
    slug = "-".join(part for part in slug.split("-") if part)
    if not slug:
        slug = "dialog"
    return f"{slug[:60]}.{extension}"


def _exported_at() -> str:
    return "Exported at " + datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


def _pdf_text(text: str) -> str:
    return "<br/>".join(escape(line) for line in text.splitlines())


def _clean_text(text: str) -> str:
    return "\n".join(line.rstrip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")).strip()
